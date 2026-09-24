"""Extração local de tabelas e texto livre de PDF e DOCX."""

import io

import pdfplumber
from docx import Document

from core.constants import LLM_MAX_BLOCK_CHARS
from core.errors import FileReadError
from core.ingest.tabular import _make_table
from core.models import RawTable, ReadMethod, TextBlock
from core.text_utils import looks_like_student_data, nfc


def _chunk_lines(lines: list[str], prefix: str) -> list[TextBlock]:
    blocks: list[TextBlock] = []
    current: list[str] = []
    start = 1
    for number, line in enumerate(lines, start=1):
        if current and sum(len(item) + 1 for item in current) + len(line) > LLM_MAX_BLOCK_CHARS:
            text = nfc("\n".join(current))
            blocks.append(
                TextBlock(f"{prefix} {start}–{number - 1}", text, looks_like_student_data([text]))
            )
            current = []
            start = number
        current.append(line)
    if current:
        text = nfc("\n".join(current))
        blocks.append(
            TextBlock(f"{prefix} {start}–{len(lines)}", text, looks_like_student_data([text]))
        )
    return blocks


def split_text_blocks(text: str, locator_prefix: str) -> list[TextBlock]:
    """Divide texto livre em blocos de até 6.000 caracteres por linha."""
    if not text.strip():
        return []
    return _chunk_lines(text.splitlines(), locator_prefix)


def read_pdf(data: bytes) -> tuple[list[RawTable], list[TextBlock]]:
    """Extrai tabelas por código e texto fora das áreas das tabelas."""
    tables: list[RawTable] = []
    blocks: list[TextBlock] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                page_tables = page.find_tables()
                outside = page
                for table_number, found in enumerate(page_tables, start=1):
                    rows = [
                        (number, ["" if cell is None else str(cell).strip() for cell in row])
                        for number, row in enumerate(found.extract(), start=1)
                    ]
                    table = _make_table(rows, f"página {page_number}, tabela {table_number}")
                    table.method = ReadMethod.TABELA_DOC
                    if table.header is not None or table.rows:
                        tables.append(table)
                    outside = outside.outside_bbox(found.bbox)
                text = nfc((outside.extract_text() or "").strip())
                if text:
                    blocks.append(
                        TextBlock(f"página {page_number}", text, looks_like_student_data([text]))
                    )
    except Exception as exc:
        raise FileReadError("ARQ_ILEGIVEL") from exc
    return tables, blocks


def read_docx(data: bytes) -> tuple[list[RawTable], list[TextBlock]]:
    """Extrai tabelas e parágrafos de um DOCX sem usar a IA."""
    try:
        document = Document(io.BytesIO(data))
        tables: list[RawTable] = []
        for table_number, found in enumerate(document.tables, start=1):
            rows = [
                (number, [cell.text.strip() for cell in row.cells])
                for number, row in enumerate(found.rows, start=1)
            ]
            table = _make_table(rows, f"tabela {table_number}")
            table.method = ReadMethod.TABELA_DOC
            if table.header is not None or table.rows:
                tables.append(table)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        return tables, _chunk_lines(paragraphs, "parágrafos") if any(paragraphs) else []
    except Exception as exc:
        raise FileReadError("ARQ_ILEGIVEL") from exc
