"""Dados fictícios e reprodutíveis para os testes."""

import re
from io import BytesIO

import openpyxl
import xlwt
from docx import Document
from fpdf import FPDF


def make_cpf(base9: str) -> str:
    """Gera CPF sintético com dígitos verificadores oficiais."""
    if re.fullmatch(r"[0-9]{9}", base9) is None or len(set(base9)) == 1:
        raise ValueError("base9 deve ter nove dígitos ASCII não repetidos")
    digits = base9
    for weight in (10, 11):
        check = (sum(int(digits[i]) * (weight - i) for i in range(weight - 1)) * 10) % 11
        digits += str(0 if check == 10 else check)
    return digits


def make_csv(
    rows: list[list[str]], encoding: str = "utf-8", sep: str = ";", header: list[str] | None = None
) -> bytes:
    """Gera CSV textual simples para os testes de leitura."""
    all_rows = ([header] if header is not None else []) + rows
    return ("\n".join(sep.join(row) for row in all_rows) + "\n").encode(encoding)


def make_xlsx(sheets: dict[str, list[list[str | int | None]]]) -> bytes:
    """Gera planilha XLSX fictícia em memória."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def make_xls(rows: list[list[str | int | None]]) -> bytes:
    """Gera planilha XLS fictícia em memória."""
    book = xlwt.Workbook()
    sheet = book.add_sheet("Alunos")
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if value is not None:
                sheet.write(row_index, column_index, value)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def make_docx(table_rows: list[list[str]], paragraphs: list[str]) -> bytes:
    """Gera DOCX com tabela e texto livre fictícios."""
    document = Document()
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for values in table_rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, values, strict=True):
                cell.text = value
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_pdf(table_rows: list[list[str]], free_text: str = "") -> bytes:
    """Gera PDF com tabela simples e texto fora da tabela."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    for row in table_rows:
        for cell in row:
            pdf.cell(70, 10, text=cell, border=1)
        pdf.ln(10)
    if free_text:
        pdf.ln(10)
        pdf.multi_cell(180, 8, text=free_text)
    return bytes(pdf.output())
